"""Knocky server: personal cabinet, guest pages, RSVP.

Runs on the Timeweb server behind Caddy (app.knocky.ru -> 127.0.0.1:8000).
Data: SQLite in $KNOCKY_DATA (default /opt/knocky/data).
Secrets (mail password, key for the sealed guest page) live only on the server
in $KNOCKY_SECRETS (default /etc/knocky/secrets.json), set via the one-time /setup page.
Nothing secret is stored in this repository.
"""
import asyncio, base64, threading, datetime as dt, hashlib, hmac, io, json, os, re, secrets, smtplib, sqlite3, time
from email.message import EmailMessage
from email.utils import formataddr
from pathlib import Path

from starlette.applications import Starlette
from starlette.middleware import Middleware
from starlette.requests import Request
from starlette.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from starlette.routing import Mount, Route
from starlette.staticfiles import StaticFiles

HERE = Path(__file__).resolve().parent
DATA = Path(os.environ.get("KNOCKY_DATA", "/opt/knocky/data"))
SECRETS = Path(os.environ.get("KNOCKY_SECRETS", "/etc/knocky/secrets.json"))
DEV = os.environ.get("KNOCKY_DEV") == "1"
MEDIA = DATA / "media"
APP_URL = os.environ.get("KNOCKY_APP_URL", "https://app.knocky.ru")
SITE_URL = "https://knocky.ru"
ORIGINS = {APP_URL, SITE_URL, "https://www.knocky.ru"} | ({"http://127.0.0.1:8765", "http://localhost:8765"} if DEV else set())
ADMINS = {"me@dudukchyan.ru", "polina.demyanenko.2004@bk.ru"}
MAIL_FROM = "hello@knocky.ru"
DOC_VERSION = "15.10.2026"
# sha256 of the one-time setup token (the token itself is given to the owner, never stored here)
SETUP_HASH = "36ece8f75ae5ca9817d22ef79eca21563bf3bd48314e6615f1623b9184150b26"
BLOCK_IDS = {"timer", "program", "map", "dress", "rsvp", "gallery", "wishlist", "named", "year", "seating", "print"}
SESSION_DAYS = 30
MAX_TABLES = 100
# prices: the same numbers as in the constructor on knocky.ru (BLOCKS, GUEST_STEP_PRICE, DAY_MAP_PRICE)
PRICE_BASE = 490
PRICE_BLOCKS = {"timer": 50, "program": 100, "map": 100, "dress": 150, "rsvp": 200, "gallery": 100, "wishlist": 100,
                "named": 300, "seating": 200, "print": 50, "year": 150}
PRICE_PACK = 200
PRICE_DAY_MAP = 50
ROBOKASSA_URL = "https://auth.robokassa.ru/Merchant/Index.aspx"
RENDER_PORT = int(os.environ.get("KNOCKY_PORT", "8000"))
RENDER_SIZES = {"og": (1200, 630)}
EMAIL_RE = re.compile(r"^[^@\s]{1,64}@[^@\s]{1,190}\.[^@\s]{2,30}$")
TR = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
              ["a","b","v","g","d","e","e","zh","z","i","y","k","l","m","n","o","p","r","s","t","u","f","h","ts","ch","sh","sch","","y","","e","yu","ya"]))
RESERVED = {"api", "i", "p", "m", "setup", "status", "login", "admin", "static", "assets", "new", "invite", "app", "www", "print", "og", "render", "pay"}


# ---------------------------------------------------------------- storage
def now() -> int:
    return int(time.time())


def db() -> sqlite3.Connection:
    c = sqlite3.connect(DATA / "knocky.db", timeout=10)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    return c


SCHEMA = """
PRAGMA journal_mode=WAL;
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, email TEXT UNIQUE NOT NULL, created INTEGER NOT NULL,
  consent_at INTEGER, consent_version TEXT, last_login INTEGER);
CREATE TABLE IF NOT EXISTS codes(email TEXT PRIMARY KEY, hash TEXT NOT NULL, expires INTEGER NOT NULL,
  attempts INTEGER NOT NULL DEFAULT 0, sent INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS sends(email TEXT, ip TEXT, at INTEGER);
CREATE TABLE IF NOT EXISTS sessions(hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  created INTEGER NOT NULL, expires INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS invitations(id TEXT PRIMARY KEY, user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
  claim TEXT, tpl TEXT NOT NULL, data TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'draft', slug TEXT UNIQUE,
  created INTEGER NOT NULL, updated INTEGER NOT NULL, published_at INTEGER, expires_at INTEGER, paid_via TEXT);
CREATE TABLE IF NOT EXISTS guests(id INTEGER PRIMARY KEY, inv_id TEXT NOT NULL REFERENCES invitations(id) ON DELETE CASCADE,
  name TEXT NOT NULL, token TEXT UNIQUE NOT NULL, pos INTEGER NOT NULL, opens INTEGER NOT NULL DEFAULT 0,
  first_open INTEGER, last_open INTEGER, active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS rsvp(id INTEGER PRIMARY KEY, inv_id TEXT NOT NULL REFERENCES invitations(id) ON DELETE CASCADE,
  guest_id INTEGER, name TEXT NOT NULL, attend INTEGER NOT NULL, people INTEGER NOT NULL DEFAULT 1, comment TEXT,
  created INTEGER NOT NULL, updated INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS wish(inv_id TEXT NOT NULL REFERENCES invitations(id) ON DELETE CASCADE, idx INTEGER NOT NULL,
  name TEXT NOT NULL, created INTEGER NOT NULL, PRIMARY KEY(inv_id, idx));
CREATE TABLE IF NOT EXISTS hits(key TEXT, at INTEGER);
CREATE TABLE IF NOT EXISTS orders(id INTEGER PRIMARY KEY AUTOINCREMENT, inv_id TEXT NOT NULL, user_id INTEGER, amount INTEGER NOT NULL,
  items TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'new', test INTEGER NOT NULL DEFAULT 0, created INTEGER NOT NULL, paid_at INTEGER);
CREATE INDEX IF NOT EXISTS inv_user ON invitations(user_id);
CREATE INDEX IF NOT EXISTS rsvp_inv ON rsvp(inv_id);
CREATE INDEX IF NOT EXISTS hits_key ON hits(key, at);
"""


def init_db():
    DATA.mkdir(parents=True, exist_ok=True)
    MEDIA.mkdir(parents=True, exist_ok=True)
    with db() as c:
        c.executescript(SCHEMA)
        # several days: which days the guest will come to (JSON list of day numbers, 0 = first)
        if "days" not in {r[1] for r in c.execute("PRAGMA table_info(rsvp)")}:
            c.execute("ALTER TABLE rsvp ADD COLUMN days TEXT")


def load_secrets() -> dict:
    try:
        return json.loads(SECRETS.read_text())
    except Exception:
        return {}


def save_secrets(d: dict):
    SECRETS.parent.mkdir(parents=True, exist_ok=True)
    tmp = SECRETS.with_suffix(".tmp")
    tmp.write_text(json.dumps(d))
    os.chmod(tmp, 0o600)
    tmp.replace(SECRETS)


def sha(s: str) -> str:
    return hashlib.sha256(s.encode()).hexdigest()


def rate(key: str, limit: int, window: int) -> bool:
    """True if allowed. Simple sliding window stored in SQLite."""
    t = now()
    with db() as c:
        c.execute("DELETE FROM hits WHERE at<?", (t - 86400,))
        n = c.execute("SELECT COUNT(*) FROM hits WHERE key=? AND at>?", (key, t - window)).fetchone()[0]
        if n >= limit:
            return False
        c.execute("INSERT INTO hits VALUES(?,?)", (key, t))
    return True


def client_ip(req: Request) -> str:
    return (req.headers.get("x-forwarded-for", "").split(",")[0].strip() or (req.client.host if req.client else "")) or "?"


# ---------------------------------------------------------------- helpers
def err(msg: str, code: int = 400):
    return JSONResponse({"error": msg}, status_code=code)


async def body(req: Request, limit: int = 64_000) -> dict:
    raw = await req.body()
    if len(raw) > limit:
        raise ValueError("Слишком большой запрос")
    try:
        d = json.loads(raw or b"{}")
    except Exception:
        raise ValueError("Неверный формат запроса")
    if not isinstance(d, dict):
        raise ValueError("Неверный формат запроса")
    return d


def user_of(req: Request):
    tok = req.cookies.get("kn")
    if not tok:
        return None
    with db() as c:
        r = c.execute("SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.hash=? AND s.expires>?",
                      (sha(tok), now())).fetchone()
    return r


def is_admin(u) -> bool:
    return bool(u) and u["email"] in ADMINS


def need_user(fn):
    async def w(req: Request):
        u = user_of(req)
        if not u:
            return err("Войдите в кабинет", 401)
        req.state.user = u
        try:
            return await fn(req)
        except ValueError as e:
            return err(str(e))
    return w


def translit(s: str) -> str:
    s = "".join(TR.get(ch, ch) for ch in s.lower())
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:40].strip("-")


def slug_ok(s: str) -> bool:
    return bool(re.fullmatch(r"[a-z0-9](?:[a-z0-9-]{1,38}[a-z0-9])", s)) and s not in RESERVED and "--" not in s


def free_slug(c, base: str, inv_id: str) -> str:
    base = translit(base) or "svadba"
    if len(base) < 3:
        base = (base + "-svadba").strip("-")
    cand = base
    for i in range(2, 200):
        if slug_ok(cand):
            r = c.execute("SELECT id FROM invitations WHERE slug=?", (cand,)).fetchone()
            if not r or r["id"] == inv_id:
                return cand
        cand = f"{base}-{i}"
    return f"{base}-{secrets.token_hex(3)}"


def clean_data(d: dict) -> dict:
    """Keep only the invitation fields the constructor uses; move photos out of the JSON."""
    if not isinstance(d, dict):
        raise ValueError("Нет данных приглашения")
    tpl = str(d.get("tpl") or "")[:30]
    if not re.fullmatch(r"[a-z0-9_-]{2,30}", tpl):
        raise ValueError("Не выбран дизайн")
    sel = [b for b in (d.get("sel") or []) if b in BLOCK_IDS]
    order = [b for b in (d.get("order") or []) if b in BLOCK_IDS]
    v = d.get("v") or {}
    if not isinstance(v, dict):
        raise ValueError("Нет данных приглашения")
    out = {}
    def s(k, n=200):
        x = v.get(k)
        if x is not None:
            out[k] = str(x)[:n]
    for k, n in (("a", 60), ("b", 60), ("date", 10), ("time", 5), ("venue", 120), ("addr", 300), ("dressNote", 300), ("rsvpBy", 10), ("d0t", 40)):
        s(k, n)
    for k in ("lat", "lng"):
        try:
            if v.get(k) not in (None, ""):
                out[k] = round(float(v[k]), 6)
        except Exception:
            pass
    if isinstance(v.get("prog"), list):
        out["prog"] = [{"t": str(r.get("t", ""))[:5], "x": str(r.get("x", ""))[:120]} for r in v["prog"][:20] if isinstance(r, dict)]
    if isinstance(v.get("sw"), list):
        out["sw"] = [x for x in (str(c)[:9] for c in v["sw"][:8]) if re.fullmatch(r"#[0-9a-fA-F]{3,8}", x)]
    if isinstance(v.get("wish"), list):
        out["wish"] = [{"n": str(w.get("n", ""))[:120], "u": str(w.get("u", ""))[:500]} for w in v["wish"][:30] if isinstance(w, dict)]
        for w in out["wish"]:
            if w["u"] and not re.match(r"^https?://", w["u"]):
                w["u"] = ""
    try:
        out["packs"] = max(0, min(9, int(v.get("packs") or 0)))
    except Exception:
        out["packs"] = 0
    if isinstance(v.get("guests"), list):
        # 100 named links in the block, more in packs of 100 (+100 приглашений)
        out["guests"] = [str(g)[:80] for g in v["guests"][:100 * (1 + out["packs"])]]
    if isinstance(v.get("dear"), list):
        # greeting per guest line: pl «Дорогие», m «Дорогой», f «Дорогая»; empty = guessed from the name
        out["dear"] = [d if d in ("pl", "m", "f") else "" for d in (str(x or "") for x in v["dear"][:len(out.get("guests", []))])]
    if isinstance(v.get("events"), list):
        ev = []
        for e in v["events"][:4]:
            if isinstance(e, dict):
                x = {"t": str(e.get("t") or "")[:40], "date": str(e.get("date") or "")[:10], "time": str(e.get("time") or "")[:5],
                     "venue": str(e.get("venue") or "")[:60], "addr": str(e.get("addr") or "")[:160],
                     "prog": [{"t": str(r.get("t", ""))[:5], "x": str(r.get("x", ""))[:120]} for r in (e.get("prog") or [])[:20] if isinstance(r, dict)],
                     "map": bool(e.get("map")), "dress": bool(e.get("dress")), "dressNote": str(e.get("dressNote") or "")[:300]}
                for k in ("lat", "lng"):
                    try:
                        if e.get(k) not in (None, ""):
                            x[k] = round(float(e[k]), 6)
                    except Exception:
                        pass
                if isinstance(e.get("sw"), list):
                    x["sw"] = [c for c in (str(c)[:9] for c in e["sw"][:8]) if re.fullmatch(r"#[0-9a-fA-F]{3,8}", c)]
                if not re.fullmatch(r"(\d{4}-\d{2}-\d{2})?", x["date"]):
                    x["date"] = ""
                if not re.fullmatch(r"(\d{2}:\d{2})?", x["time"]):
                    x["time"] = ""
                ev.append(x)
        out["events"] = ev
    tables, seat = clean_seating(v, len(out.get("guests", [])))
    out["tables"], out["seat"] = tables, seat
    if isinstance(v.get("pal"), dict):
        try:
            out["pal"] = {"id": str(v["pal"].get("id", tpl))[:30], "i": max(0, min(5, int(v["pal"].get("i", 0))))}
        except Exception:
            pass
    photos = []
    for p in (v.get("photos") or [])[:12]:
        p = str(p)
        m = re.match(r"^data:image/(jpeg|png|webp);base64,(.+)$", p, re.S)
        if m:
            raw = base64.b64decode(m.group(2), validate=False)
            if len(raw) > 2_500_000:
                raise ValueError("Фото слишком большое")
            ext = {"jpeg": "jpg", "png": "png", "webp": "webp"}[m.group(1)]
            name = hashlib.sha256(raw).hexdigest()[:32] + "." + ext
            f = MEDIA / name
            if not f.exists():
                f.write_bytes(raw)
            photos.append(f"{APP_URL}/m/{name}")
        elif re.match(r"^" + re.escape(APP_URL) + r"/m/[0-9a-f]{32}\.(jpg|png|webp)$", p):
            photos.append(p)
    out["photos"] = photos
    return {"tpl": tpl, "sel": sel, "order": order, "v": out}


def clean_seating(v: dict, n_guests: int):
    """Tables [{n: name or number, s: seats}] and seat[i] = table index for guest line i (None = no table)."""
    tables = []
    for t in (v.get("tables") or [])[:MAX_TABLES]:
        if isinstance(t, dict):
            try:
                seats = max(0, min(100, int(t.get("s") or 0)))
            except Exception:
                seats = 0
            tables.append({"n": str(t.get("n") or "")[:20], "s": seats})
    seat = []
    for x in (v.get("seat") or [])[:n_guests]:
        try:
            k = int(x)
        except Exception:
            k = -1
        seat.append(k if 0 <= k < len(tables) else None)
    return tables, seat


def table_label(v: dict, k) -> str:
    t = v.get("tables") or []
    if k is None or not (0 <= k < len(t)):
        return ""
    return (t[k].get("n") or "").strip() or str(k + 1)


def guest_tables(d: dict) -> dict:
    """guest name -> table label (only when the seating block is bought)."""
    if "seating" not in d.get("sel", []) or "named" not in d.get("sel", []):
        return {}
    v = d["v"]
    seat = v.get("seat") or []
    out = {}
    for i, g in enumerate(v.get("guests") or []):
        g = str(g).strip()
        if g and g not in out and i < len(seat):
            lab = table_label(v, seat[i])
            if lab:
                out[g] = lab
    return out


def inv_title(d: dict) -> str:
    v = d.get("v", {})
    a, b = (v.get("a") or "").strip(), (v.get("b") or "").strip()
    return f"{a} и {b}" if a or b else "Без имен"


def inv_json(r, c, full=True) -> dict:
    d = json.loads(r["data"])
    o = {"id": r["id"], "title": inv_title(d), "tpl": r["tpl"], "status": r["status"], "slug": r["slug"],
         "created": r["created"], "updated": r["updated"], "published_at": r["published_at"], "expires_at": r["expires_at"],
         "paid_via": r["paid_via"], "date": d["v"].get("date"), "sel": d["sel"],
         "link": f"{SITE_URL}/{r['slug']}" if r["slug"] else None,
         "expired": bool(r["expires_at"] and r["expires_at"] < now())}
    st = c.execute("SELECT COUNT(*) n, SUM(attend) yes, SUM(CASE WHEN attend=1 THEN people ELSE 0 END) ppl FROM rsvp WHERE inv_id=?", (r["id"],)).fetchone()
    o["answers"] = {"total": st["n"] or 0, "yes": st["yes"] or 0, "no": (st["n"] or 0) - (st["yes"] or 0), "people": st["ppl"] or 0}
    if full:
        o["data"] = d
    return o


def own_inv(c, req, inv_id):
    u = req.state.user
    r = c.execute("SELECT * FROM invitations WHERE id=?", (inv_id,)).fetchone()
    if not r or (r["user_id"] != u["id"] and not is_admin(u)):
        return None
    return r


def sync_guests(c, inv_id: str, d: dict):
    """One named link per guest line; existing names keep their link."""
    names = [g.strip() for g in d["v"].get("guests", []) if g.strip()] if "named" in d["sel"] else []
    old = {g["name"]: g for g in c.execute("SELECT * FROM guests WHERE inv_id=?", (inv_id,))}
    keep = set()
    for i, n in enumerate(names):
        if n in old and old[n]["id"] not in keep:
            c.execute("UPDATE guests SET pos=?, active=1 WHERE id=?", (i, old[n]["id"]))
            keep.add(old[n]["id"])
        else:
            cur = c.execute("INSERT INTO guests(inv_id,name,token,pos) VALUES(?,?,?,?)", (inv_id, n, secrets.token_urlsafe(6), i))
            keep.add(cur.lastrowid)
    for g in old.values():
        if g["id"] not in keep:
            c.execute("UPDATE guests SET active=0 WHERE id=?", (g["id"],))


def add_months(ts: int, months: int) -> int:
    d = dt.datetime.fromtimestamp(ts)
    m = d.month - 1 + months
    y, m = d.year + m // 12, m % 12 + 1
    day = min(d.day, [31, 29 if y % 4 == 0 and (y % 100 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1])
    return int(d.replace(year=y, month=m, day=day).timestamp())


# ---------------------------------------------------------------- mail
def send_mail(to: str, subject: str, text: str, html: str, creds: dict | None = None):
    s = creds or load_secrets()
    pw = s.get("smtp_password")
    if DEV and not pw:
        (DATA / "outbox.txt").open("a").write(f"--- {to}\n{subject}\n{text}\n")
        return
    if not pw:
        raise RuntimeError("Почта сервера еще не настроена")
    m = EmailMessage()
    m["From"] = formataddr(("Knocky", MAIL_FROM))
    m["To"] = to
    m["Subject"] = subject
    m.set_content(text)
    m.add_alternative(html, subtype="html")
    with smtplib.SMTP_SSL("smtp.yandex.ru", 465, timeout=20) as srv:
        srv.login(s.get("smtp_login", MAIL_FROM), pw)
        srv.send_message(m)


def mail_html(title: str, body_html: str) -> str:
    return f"""<!doctype html><html><body style="margin:0;background:#F3EEEE;font-family:Arial,Helvetica,sans-serif;color:#241B20">
<table width="100%" cellpadding="0" cellspacing="0" style="background:#F3EEEE;padding:32px 12px"><tr><td align="center">
<table width="100%" cellpadding="0" cellspacing="0" style="max-width:480px;background:#FFFFFF;border-radius:16px;padding:32px">
<tr><td style="font-size:22px;font-weight:bold;color:#C62828;padding-bottom:20px">knocky</td></tr>
<tr><td style="font-size:20px;font-weight:bold;padding-bottom:12px">{title}</td></tr>
<tr><td style="font-size:16px;line-height:1.5">{body_html}</td></tr>
<tr><td style="font-size:13px;color:#6F6068;padding-top:28px">Knocky - приглашения на свадьбу по ссылке. Вопросы: hello@knocky.ru</td></tr>
</table></td></tr></table></body></html>"""


# ---------------------------------------------------------------- auth
async def auth_request(req: Request):
    try:
        d = await body(req)
    except ValueError as e:
        return err(str(e))
    email = str(d.get("email", "")).strip().lower()
    if not EMAIL_RE.match(email):
        return err("Введите почту, например anna@mail.ru")
    with db() as c:
        known = c.execute("SELECT consent_at FROM users WHERE email=?", (email,)).fetchone()
    if not (known and known["consent_at"]) and not d.get("agree"):
        return JSONResponse({"error": "Нужно согласие на обработку персональных данных", "need_agree": True}, status_code=400)
    ip = client_ip(req)
    if not rate("ip:" + ip, 10, 3600) or not rate("mail:" + email, 5, 3600):
        return err("Слишком много попыток. Попробуйте через час", 429)
    with db() as c:
        r = c.execute("SELECT sent FROM codes WHERE email=?", (email,)).fetchone()
        if r and r["sent"] > now() - 50:
            return JSONResponse({"error": "Код уже отправлен. Новый можно запросить через минуту", "already": True,
                                 "wait": 50 - (now() - r["sent"])}, status_code=429)
    code = f"{secrets.randbelow(1_000_000):06d}"
    try:
        await asyncio.to_thread(send_mail, email, f"Код для входа в Knocky: {code}",
                                f"Ваш код для входа в Knocky: {code}\nКод действует 10 минут. Если вы не запрашивали код, просто удалите письмо.",
                                mail_html("Код для входа", f'<div style="font-size:34px;font-weight:bold;letter-spacing:6px;padding:8px 0 16px">{code}</div>'
                                          "Введите его на странице входа. Код действует 10 минут.<br><br>"
                                          '<span style="color:#6F6068">Если вы не запрашивали код, просто удалите письмо.</span>'))
    except Exception as e:
        print("mail error:", repr(e), flush=True)
        return err("Не получилось отправить письмо. Попробуйте позже или напишите на hello@knocky.ru", 503)
    with db() as c:
        c.execute("INSERT OR REPLACE INTO codes VALUES(?,?,?,0,?)", (email, sha(email + ":" + code), now() + 600, now()))
        if d.get("agree"):
            c.execute("INSERT INTO users(email,created,consent_at,consent_version) VALUES(?,?,?,?) "
                      "ON CONFLICT(email) DO UPDATE SET consent_at=COALESCE(consent_at,excluded.consent_at),"
                      "consent_version=COALESCE(consent_version,excluded.consent_version)", (email, now(), now(), DOC_VERSION))
    return JSONResponse({"ok": True})


async def auth_verify(req: Request):
    try:
        d = await body(req)
    except ValueError as e:
        return err(str(e))
    email = str(d.get("email", "")).strip().lower()
    code = re.sub(r"\D", "", str(d.get("code", "")))
    with db() as c:
        r = c.execute("SELECT * FROM codes WHERE email=?", (email,)).fetchone()
        if not r or r["expires"] < now():
            return err("Код устарел. Запросите новый")
        if r["attempts"] >= 5:
            return err("Слишком много неверных попыток. Запросите новый код")
        if not hmac.compare_digest(r["hash"], sha(email + ":" + code)):
            c.execute("UPDATE codes SET attempts=attempts+1 WHERE email=?", (email,))
            return err("Неверный код")
        c.execute("DELETE FROM codes WHERE email=?", (email,))
        c.execute("INSERT INTO users(email,created) VALUES(?,?) ON CONFLICT(email) DO NOTHING", (email, now()))
        u = c.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
        c.execute("UPDATE users SET last_login=? WHERE id=?", (now(), u["id"]))
        tok = secrets.token_urlsafe(32)
        c.execute("INSERT INTO sessions VALUES(?,?,?,?)", (sha(tok), u["id"], now(), now() + SESSION_DAYS * 86400))
        c.execute("DELETE FROM sessions WHERE expires<?", (now(),))
        claim = str(d.get("claim") or "")
        if claim:
            c.execute("UPDATE invitations SET user_id=?, claim=NULL WHERE claim=? AND user_id IS NULL", (u["id"], claim))
    res = JSONResponse({"ok": True, "user": {"email": u["email"], "admin": u["email"] in ADMINS}})
    res.set_cookie("kn", tok, max_age=SESSION_DAYS * 86400, httponly=True, secure=not DEV, samesite="lax", path="/")
    return res


def new_session(c, uid: int) -> str:
    tok = secrets.token_urlsafe(32)
    c.execute("INSERT INTO sessions VALUES(?,?,?,?)", (sha(tok), uid, now(), now() + SESSION_DAYS * 86400))
    return tok


async def auth_password(req: Request):
    """Administrators can sign in with a password set on the /setup page (no email needed)."""
    try:
        d = await body(req)
    except ValueError as e:
        return err(str(e))
    email = str(d.get("email", "")).strip().lower()
    pw = str(d.get("password", ""))
    if not rate("pwip:" + client_ip(req), 10, 3600) or not rate("pw:" + email, 10, 3600):
        return err("Слишком много попыток. Попробуйте через час", 429)
    stored = load_secrets().get("admin_pw", {}).get(email)
    if not stored or email not in ADMINS or not pw_ok(pw, stored):
        return err("Неверная почта или пароль")
    with db() as c:
        c.execute("INSERT INTO users(email,created) VALUES(?,?) ON CONFLICT(email) DO NOTHING", (email, now()))
        u = c.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
        c.execute("UPDATE users SET last_login=? WHERE id=?", (now(), u["id"]))
        tok = new_session(c, u["id"])
    res = JSONResponse({"ok": True, "user": {"email": email, "admin": True}})
    res.set_cookie("kn", tok, max_age=SESSION_DAYS * 86400, httponly=True, secure=not DEV, samesite="lax", path="/")
    return res


async def auth_logout(req: Request):
    tok = req.cookies.get("kn")
    if tok:
        with db() as c:
            c.execute("DELETE FROM sessions WHERE hash=?", (sha(tok),))
    res = JSONResponse({"ok": True})
    res.delete_cookie("kn", path="/")
    return res


async def me(req: Request):
    u = user_of(req)
    if not u:
        return JSONResponse({"user": None})
    return JSONResponse({"user": {"email": u["email"], "admin": is_admin(u)}, "mail_ready": bool(load_secrets().get("smtp_password"))})


# ---------------------------------------------------------------- invitations
@need_user
async def inv_list(req: Request):
    u = req.state.user
    every = req.query_params.get("all") == "1" and is_admin(u)
    with db() as c:
        if every:
            rows = c.execute("SELECT i.*, u.email owner FROM invitations i LEFT JOIN users u ON u.id=i.user_id ORDER BY i.updated DESC LIMIT 500").fetchall()
        else:
            rows = c.execute("SELECT i.*, NULL owner FROM invitations i WHERE user_id=? ORDER BY updated DESC", (u["id"],)).fetchall()
        out = []
        for r in rows:
            o = inv_json(r, c, full=False)
            if every:
                o["owner"] = r["owner"]
            out.append(o)
    return JSONResponse({"items": out})


@need_user
async def inv_create(req: Request):
    d = clean_data((await body(req, 40_000_000)).get("data"))
    iid = secrets.token_urlsafe(8)
    with db() as c:
        c.execute("INSERT INTO invitations(id,user_id,tpl,data,created,updated) VALUES(?,?,?,?,?,?)",
                  (iid, req.state.user["id"], d["tpl"], json.dumps(d, ensure_ascii=False), now(), now()))
        sync_guests(c, iid, d)
        return JSONResponse(inv_json(c.execute("SELECT * FROM invitations WHERE id=?", (iid,)).fetchone(), c))


async def draft_create(req: Request):
    """Anonymous draft from the constructor on knocky.ru; claimed after login."""
    if not rate("draft:" + client_ip(req), 30, 3600):
        return err("Слишком много попыток", 429)
    try:
        d = clean_data((await body(req, 40_000_000)).get("data"))
    except ValueError as e:
        return err(str(e))
    u = user_of(req)
    iid, claim = secrets.token_urlsafe(8), (None if u else secrets.token_urlsafe(18))
    with db() as c:
        c.execute("INSERT INTO invitations(id,user_id,claim,tpl,data,created,updated) VALUES(?,?,?,?,?,?,?)",
                  (iid, u["id"] if u else None, claim, d["tpl"], json.dumps(d, ensure_ascii=False), now(), now()))
        sync_guests(c, iid, d)
    return JSONResponse({"id": iid, "claim": claim, "next": f"{APP_URL}/?claim={claim}" if claim else f"{APP_URL}/inv/{iid}"})


@need_user
async def draft_claim(req: Request):
    claim = str((await body(req)).get("claim") or "")[:40]
    with db() as c:
        r = c.execute("SELECT id FROM invitations WHERE claim=? AND user_id IS NULL", (claim,)).fetchone() if claim else None
        if not r:
            return err("Черновик не найден", 404)
        c.execute("UPDATE invitations SET user_id=?, claim=NULL WHERE id=?", (req.state.user["id"], r["id"]))
    return JSONResponse({"id": r["id"]})


@need_user
async def inv_get(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        return JSONResponse(inv_json(r, c))


@need_user
async def inv_update(req: Request):
    b = await body(req, 40_000_000)
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        if "data" in b:
            nd = clean_data(b["data"])
            if r["status"] == "published":
                # after payment the set of blocks is fixed; texts, time and address can change
                nd["sel"] = d["sel"]
            d = nd
        slug = r["slug"]
        if "slug" in b and b["slug"] != r["slug"]:
            want = str(b["slug"] or "").strip().lower()
            if not slug_ok(want):
                raise ValueError("Адрес: от 3 до 40 латинских букв, цифр и дефисов")
            if c.execute("SELECT 1 FROM invitations WHERE slug=? AND id<>?", (want, r["id"])).fetchone():
                raise ValueError("Этот адрес уже занят, выберите другой")
            if r["status"] == "published" and not is_admin(req.state.user):
                raise ValueError("Адрес опубликованного приглашения уже у гостей, его нельзя поменять")
            slug = want
        c.execute("UPDATE invitations SET tpl=?, data=?, slug=?, updated=? WHERE id=?",
                  (d["tpl"], json.dumps(d, ensure_ascii=False), slug, now(), r["id"]))
        sync_guests(c, r["id"], d)
        return JSONResponse(inv_json(c.execute("SELECT * FROM invitations WHERE id=?", (r["id"],)).fetchone(), c))


@need_user
async def inv_delete(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        if r["status"] == "published" and not is_admin(req.state.user):
            return err("Опубликованное приглашение удалить нельзя: ссылка уже у гостей. Напишите на hello@knocky.ru")
        c.execute("DELETE FROM invitations WHERE id=?", (r["id"],))
    return JSONResponse({"ok": True})


@need_user
async def inv_publish(req: Request):
    """Admin publishes without payment. Paid publishing arrives with Robokassa."""
    u = req.state.user
    if not is_admin(u):
        return err("Оплата появится в ближайшее время. Пока опубликовать приглашение может администратор", 402)
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        slug = r["slug"] or free_slug(c, f"{d['v'].get('a','')}-{d['v'].get('b','')}", r["id"])
        t = now()
        exp = add_months(t, 12 if "year" in d["sel"] else 3)
        c.execute("UPDATE invitations SET status='published', slug=?, published_at=COALESCE(published_at,?), expires_at=?, paid_via=COALESCE(paid_via,'admin'), updated=? WHERE id=?",
                  (slug, t, exp, t, r["id"]))
        return JSONResponse(inv_json(c.execute("SELECT * FROM invitations WHERE id=?", (r["id"],)).fetchone(), c))


@need_user
async def inv_unpublish(req: Request):
    if not is_admin(req.state.user):
        return err("Нет доступа", 403)
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        c.execute("UPDATE invitations SET status='draft', updated=? WHERE id=?", (now(), r["id"]))
        return JSONResponse(inv_json(c.execute("SELECT * FROM invitations WHERE id=?", (r["id"],)).fetchone(), c))


def answers_of(c, inv_id):
    rows = c.execute("SELECT r.*, g.name gname FROM rsvp r LEFT JOIN guests g ON g.id=r.guest_id WHERE r.inv_id=? ORDER BY r.updated DESC", (inv_id,)).fetchall()
    def days(x):
        try:
            return json.loads(x) if x else None
        except Exception:
            return None
    return [{"name": r["name"], "guest": r["gname"], "attend": bool(r["attend"]), "people": r["people"], "comment": r["comment"] or "", "at": r["updated"],
             "days": days(r["days"])} for r in rows]


DAY_N = ["первый", "второй", "третий", "четвертый", "пятый"]


def day_titles(v: dict) -> list:
    """Tab names the couple typed, or «День первый», «День второй»..."""
    ev = v.get("events") or []
    out = []
    for i in range(len(ev) + 1):
        t = (v.get("d0t") if i == 0 else (ev[i - 1] or {}).get("t")) or ""
        out.append(t.strip() or f"День {DAY_N[i]}")
    return out


def guests_of(c, inv_id, slug, d: dict | None = None):
    rows = c.execute("SELECT g.*, (SELECT attend FROM rsvp WHERE guest_id=g.id ORDER BY updated DESC LIMIT 1) ans, "
                     "(SELECT days FROM rsvp WHERE guest_id=g.id ORDER BY updated DESC LIMIT 1) ans_days "
                     "FROM guests g WHERE inv_id=? AND active=1 ORDER BY pos", (inv_id,)).fetchall()
    tb = guest_tables(d) if d else {}
    return [{"name": g["name"], "link": f"{SITE_URL}/{slug}?g={g['token']}" if slug else None, "opens": g["opens"],
             "first_open": g["first_open"], "answer": None if g["ans"] is None else bool(g["ans"]), "table": tb.get(g["name"]),
             "days": json.loads(g["ans_days"]) if g["ans_days"] else None} for g in rows]


@need_user
async def inv_answers(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        wl = json.loads(r["data"])["v"].get("wish", [])
        taken = {w["idx"]: w["name"] for w in c.execute("SELECT * FROM wish WHERE inv_id=?", (r["id"],))}
        return JSONResponse({"answers": answers_of(c, r["id"]), "guests": guests_of(c, r["id"], r["slug"], json.loads(r["data"])),
                             "days": day_titles(json.loads(r["data"])["v"]),
                             "wish": [{"n": w["n"], "by": taken.get(i)} for i, w in enumerate(wl) if w.get("n")]})


@need_user
async def inv_seating(req: Request):
    """Seating from the cabinet: tables and who sits where (also editable after publishing)."""
    b = await body(req)
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        if "seating" not in d["sel"]:
            return err("Рассадка не подключена в этом приглашении")
        tables, seat = clean_seating(b, len(d["v"].get("guests", [])))
        d["v"]["tables"], d["v"]["seat"] = tables, seat
        c.execute("UPDATE invitations SET data=?, updated=? WHERE id=?", (json.dumps(d, ensure_ascii=False), now(), r["id"]))
        return JSONResponse({"tables": tables, "seat": seat, "guests": guests_of(c, r["id"], r["slug"], d)})


@need_user
async def inv_print(req: Request):
    """Data for the printable version: design, texts, one card per named guest with his link."""
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        if "print" not in d["sel"]:
            return err("Печатная версия не подключена в этом приглашении", 403)
        v = dict(d["v"])
        raw = [str(x) for x in v.pop("guests", [])]
        dear = v.pop("dear", [])
        for k in ("photos", "wish", "tables", "seat"):
            v.pop(k, None)
        cards = []
        link = f"{SITE_URL}/{r['slug']}" if r["slug"] and r["status"] == "published" else None
        if "named" in d["sel"]:
            for g in guests_of(c, r["id"], r["slug"], d):
                idx = next((i for i, x in enumerate(raw) if x.strip() == g["name"]), -1)
                cards.append({"guest": g["name"], "dear": dear[idx] if 0 <= idx < len(dear) else "",
                              "link": g["link"] if link else None, "table": g["table"]})
        if not cards:
            cards.append({"guest": None, "dear": "", "link": link, "table": None})
    return JSONResponse({"tpl": d["tpl"], "sel": d["sel"], "order": d["order"], "v": v, "cards": cards,
                         "published": bool(link)}, headers={"Cache-Control": "no-store"})


@need_user
async def inv_export(req: Request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        ans, gs = answers_of(c, r["id"]), guests_of(c, r["id"], r["slug"], json.loads(r["data"]))
        dt_ = day_titles(json.loads(r["data"])["v"])
        title = inv_title(json.loads(r["data"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Ответы"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="C62828")
    def fmt(ts):
        return dt.datetime.fromtimestamp(ts).strftime("%d.%m.%Y %H:%M") if ts else ""
    multi = len(dt_) > 1

    def per_day(attend, days):
        """One column per day: Придет / Не придет (only when the invitation has several days)."""
        if not multi:
            return []
        if attend is None:
            return [""] * len(dt_)
        return ["Придет" if attend and (days is None or i in days) else "Не придет" for i in range(len(dt_))]
    ws.append(["Гость", "Ответ", "Сколько человек"] + dt_ * multi + ["Комментарий", "Именная ссылка", "Когда ответил"])
    for a in ans:
        ws.append([a["name"], "Придет" if a["attend"] else "Не сможет", a["people"] if a["attend"] else 0] + per_day(bool(a["attend"]), a["days"])
                  + [a["comment"], a["guest"] or "", fmt(a["at"])])
    ws.append([])
    ws.append(["Итого придут", "", sum(a["people"] for a in ans if a["attend"])])
    if multi:
        for i, t in enumerate(dt_):
            ws.append([t, "", sum(a["people"] for a in ans if a["attend"] and (a["days"] is None or i in a["days"]))])
    if gs:
        g2 = wb.create_sheet("Именные ссылки")
        g2.append(["Гость", "Ссылка", "Открывал", "Первое открытие", "Ответ"] + dt_ * multi + ["Стол"])
        for g in gs:
            g2.append([g["name"], g["link"] or "", g["opens"], fmt(g["first_open"]), "" if g["answer"] is None else ("Придет" if g["answer"] else "Не сможет")]
                      + per_day(g["answer"], g["days"]) + [g["table"] or ""])
    # guest texts like "=)" must stay text (not a formula), and empty texts must be empty cells, otherwise Excel offers to "repair" the file
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    txt = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                    # an empty text cell is written as <c t="inlineStr"></c>, which Excel calls broken: leave it truly empty
                    cell.value = txt or None
                    if txt:
                        cell.data_type = "s"
    for sh in wb.worksheets:
        for cell in sh[1]:
            cell.font, cell.fill = head, fill
        for col in sh.columns:
            w = max(len(str(x.value or "")) for x in col)
            sh.column_dimensions[col[0].column_letter].width = min(60, max(12, w + 2))
    buf = io.BytesIO()
    wb.save(buf)
    fn = "knocky-" + (translit(title) or "otvety") + ".xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fn}"'})


# ---------------------------------------------------------------- payment (Robokassa)
def price_of(d: dict) -> tuple[int, list]:
    """Price the server trusts: base + blocks + guest packs + maps for extra days."""
    sel = set(d["sel"])
    v = d["v"]
    items = [("Основа приглашения", PRICE_BASE)]
    names = {"timer": "Таймер", "program": "Программа дня", "map": "Карта и маршрут", "dress": "Дресс-код", "rsvp": "Анкета гостя",
             "gallery": "Фотогалерея", "wishlist": "Вишлист", "named": "Именные ссылки", "seating": "Рассадка",
             "print": "Печатная версия", "year": "Ссылка на 12 месяцев"}
    for b, pr in PRICE_BLOCKS.items():
        if b in sel and (b != "seating" or "named" in sel):
            items.append((names[b], pr))
    if "named" in sel and int(v.get("packs") or 0):
        n = int(v["packs"])
        items.append((f"+100 приглашений x {n}", n * PRICE_PACK))
    if "map" in sel:
        n = sum(1 for e in (v.get("events") or []) if e.get("map"))
        if n:
            items.append((f"Карта для дополнительного дня x {n}", n * PRICE_DAY_MAP))
    return sum(x for _, x in items), items


def rk_conf() -> dict | None:
    s = load_secrets().get("robokassa") or {}
    test = bool(s.get("test", True))
    p1, p2 = (s.get("test_p1"), s.get("test_p2")) if test else (s.get("p1"), s.get("p2"))
    if not s.get("login") or not p1 or not p2:
        return None
    return {"login": s["login"], "p1": p1, "p2": p2, "test": test, "sno": s.get("sno") or "usn_income"}


def rk_sign(*parts) -> str:
    return hashlib.sha256(":".join(str(x) for x in parts).encode()).hexdigest()


@need_user
async def inv_price(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        total, items = price_of(json.loads(r["data"]))
    conf = rk_conf()
    # while Robokassa is in test mode only admins can pay (a test payment would publish for free)
    ready = conf is not None and (not conf["test"] or is_admin(req.state.user))
    return JSONResponse({"total": total, "items": [{"name": n, "price": p} for n, p in items], "ready": ready,
                         "test": bool(conf and conf["test"])})


@need_user
async def inv_pay(req: Request):
    """Create an order and send the client to Robokassa. The invitation is published when Robokassa confirms the payment."""
    from urllib.parse import quote, urlencode
    conf = rk_conf()
    u = req.state.user
    if not conf or (conf["test"] and not is_admin(u)):
        return err("Прием оплаты откроется в ближайшие дни. Приглашение сохранено в кабинете", 503)
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        if r["status"] == "published" and not (r["expires_at"] and r["expires_at"] < now()):
            return err("Приглашение уже опубликовано")
        d = json.loads(r["data"])
        total, items = price_of(d)
        cur = c.execute("INSERT INTO orders(inv_id,user_id,amount,items,test,created) VALUES(?,?,?,?,?,?)",
                        (r["id"], u["id"], total, json.dumps(items, ensure_ascii=False), 1 if conf["test"] else 0, now()))
        oid = cur.lastrowid
    out_sum = f"{total}.00"
    receipt = json.dumps({"sno": conf["sno"], "items": [{"name": "Онлайн-приглашение Knocky: " + inv_title(d)[:80], "quantity": 1, "sum": total,
                          "payment_method": "full_payment", "payment_object": "service", "tax": "none"}]}, ensure_ascii=False, separators=(",", ":"))
    rq = quote(receipt, safe="")
    sig = rk_sign(conf["login"], out_sum, oid, rq, conf["p1"])
    params = {"MerchantLogin": conf["login"], "OutSum": out_sum, "InvId": oid, "Description": "Приглашение Knocky",
              "Receipt": rq, "SignatureValue": sig, "Email": u["email"], "Culture": "ru", "Encoding": "utf-8"}
    if conf["test"]:
        params["IsTest"] = 1
    # Receipt is already encoded once (it is signed in that form); urlencode encodes it a second time, as Robokassa expects in a link
    return JSONResponse({"url": ROBOKASSA_URL + "?" + urlencode(params), "amount": total, "order": oid})


def publish_paid(c, inv_id: str, via: str):
    r = c.execute("SELECT * FROM invitations WHERE id=?", (inv_id,)).fetchone()
    if not r:
        return
    d = json.loads(r["data"])
    slug = r["slug"] or free_slug(c, f"{d['v'].get('a','')}-{d['v'].get('b','')}", r["id"])
    t = now()
    exp = add_months(t, 12 if "year" in d["sel"] else 3)
    c.execute("UPDATE invitations SET status='published', slug=?, published_at=COALESCE(published_at,?), expires_at=?, paid_via=?, updated=? WHERE id=?",
              (slug, t, exp, via, t, r["id"]))


async def pay_result(req: Request):
    """Robokassa Result URL: confirms the payment (server to server). Answer OK<InvId>."""
    form = dict(await req.form()) if req.method == "POST" else dict(req.query_params)
    try:
        oid = int(form.get("InvId", 0))
        out_sum = str(form.get("OutSum", ""))
        sig = str(form.get("SignatureValue", "")).lower()
    except Exception:
        return Response("bad request", status_code=400)
    s = load_secrets().get("robokassa") or {}
    with db() as c:
        o = c.execute("SELECT * FROM orders WHERE id=?", (oid,)).fetchone()
        if not o:
            return Response("bad order", status_code=400)
        p2 = s.get("test_p2") if o["test"] else s.get("p2")
        if not p2 or not hmac.compare_digest(rk_sign(out_sum, oid, p2), sig):
            return Response("bad sign", status_code=400)
        try:
            if round(float(out_sum)) != o["amount"]:
                return Response("bad sum", status_code=400)
        except Exception:
            return Response("bad sum", status_code=400)
        if o["status"] != "paid":
            c.execute("UPDATE orders SET status='paid', paid_at=? WHERE id=?", (now(), oid))
            publish_paid(c, o["inv_id"], "robokassa-test" if o["test"] else "robokassa")
    return Response(f"OK{oid}", media_type="text/plain")


async def pay_back(req: Request):
    """Success / Fail URL: back to the invitation in the cabinet."""
    ok = req.path_params["kind"] == "ok"
    try:
        oid = int(req.query_params.get("InvId", 0))
    except Exception:
        oid = 0
    with db() as c:
        o = c.execute("SELECT inv_id FROM orders WHERE id=?", (oid,)).fetchone()
    target = f"/inv/{o['inv_id']}" if o else "/"
    return RedirectResponse(target + ("?paid=1" if ok else "?payfail=1"), status_code=302)


# ---------------------------------------------------------------- public (guests)
def public_payload(c, r, token: str = ""):
    d = json.loads(r["data"])
    v = dict(d["v"])
    sel = list(d["sel"])
    guest = None
    if token:
        g = c.execute("SELECT * FROM guests WHERE inv_id=? AND token=? AND active=1", (r["id"], token)).fetchone()
        if g:
            guest = g["name"]
    for k in ("guests", "dear", "tables", "seat"):
        v.pop(k, None)
    if guest and "named" in sel:
        raw = [str(x) for x in d["v"].get("guests", [])]
        idx = next((i for i, x in enumerate(raw) if x.strip() == guest), -1)
        dear = d["v"].get("dear", [])
        v["guests"] = [guest]
        v["dear"] = [dear[idx] if 0 <= idx < len(dear) else ""]
        lab = guest_tables(d).get(guest)
        if lab:
            v["tables"], v["seat"] = [{"n": lab, "s": 0}], [0]
    else:
        sel = [x for x in sel if x not in ("named", "seating")]
    taken = [w["idx"] for w in c.execute("SELECT idx FROM wish WHERE inv_id=?", (r["id"],))]
    return {"tpl": d["tpl"], "sel": sel, "order": d["order"], "v": v, "guest": guest, "wishTaken": taken,
            "status": r["status"], "expired": bool(r["expires_at"] and r["expires_at"] < now()), "slug": r["slug"]}


async def pub_get(req: Request):
    slug = req.path_params["slug"].lower()
    tok = req.query_params.get("g", "")[:20]
    with db() as c:
        r = c.execute("SELECT * FROM invitations WHERE slug=? AND status='published'", (slug,)).fetchone()
        if not r:
            return err("Приглашение не найдено", 404)
        p = public_payload(c, r, tok)
        if p["expired"]:
            return JSONResponse({"expired": True, "v": {"a": p["v"].get("a"), "b": p["v"].get("b")}})
        if tok and p["guest"]:
            c.execute("UPDATE guests SET opens=opens+1, first_open=COALESCE(first_open,?), last_open=? WHERE inv_id=? AND token=?", (now(), now(), r["id"], tok))
    return JSONResponse(p, headers={"Cache-Control": "no-store"})


@need_user
async def preview_get(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        p = public_payload(c, r)
        names = [g for g in d["v"].get("guests", []) if g.strip()]
        if "named" in d["sel"]:
            p["sel"] = d["sel"]
            p["v"]["guests"] = names[:1] or ["Ольга и Сергей"]
            raw = d["v"].get("guests", [])
            first = next((i for i, x in enumerate(raw) if str(x).strip()), -1)
            dear = d["v"].get("dear", [])
            p["v"]["dear"] = [dear[first] if 0 <= first < len(dear) else ""]
            lab = guest_tables(d).get(names[0].strip()) if names else None
            if lab:
                p["v"]["tables"], p["v"]["seat"] = [{"n": lab, "s": 0}], [0]
        p["preview"] = True
    return JSONResponse(p)


async def pub_rsvp(req: Request):
    slug = req.path_params["slug"].lower()
    if not rate("rsvp:" + client_ip(req), 30, 3600):
        return err("Слишком много ответов с одного устройства. Попробуйте позже", 429)
    try:
        d = await body(req)
    except ValueError as e:
        return err(str(e))
    name = re.sub(r"\s+", " ", str(d.get("name", ""))).strip()[:80]
    if not name:
        return err("Укажите имя")
    attend = 1 if d.get("attend") else 0
    try:
        people = max(1, min(20, int(d.get("people") or 1)))
    except Exception:
        people = 1
    comment = str(d.get("comment") or "")[:500]
    tok = str(d.get("g") or "")[:20]
    days_in = d.get("days")
    with db() as c:
        r = c.execute("SELECT * FROM invitations WHERE slug=? AND status='published'", (slug,)).fetchone()
        if not r or (r["expires_at"] and r["expires_at"] < now()):
            return err("Приглашение не найдено", 404)
        dd = json.loads(r["data"])
        if "rsvp" not in dd["sel"]:
            return err("В этом приглашении нет анкеты")
        n_days = len(dd["v"].get("events") or []) + 1
        days = None
        if n_days > 1 and attend:
            try:
                days = sorted({int(x) for x in (days_in or []) if 0 <= int(x) < n_days})
            except Exception:
                days = []
            if not days:
                return err("Отметьте, в какие дни вы придете")
        g = c.execute("SELECT id FROM guests WHERE inv_id=? AND token=?", (r["id"], tok)).fetchone() if tok else None
        prev = c.execute("SELECT id FROM rsvp WHERE inv_id=? AND guest_id=?", (r["id"], g["id"])).fetchone() if g else None
        if not g:
            # the same guest answering again from the common link: keep only the latest answer (matched by name)
            key = name.casefold().replace("ё", "е")
            prev = next((x for x in c.execute("SELECT id, name FROM rsvp WHERE inv_id=? AND guest_id IS NULL ORDER BY id DESC", (r["id"],))
                         if x["name"].casefold().replace("ё", "е") == key), None)
        if prev:
            c.execute("UPDATE rsvp SET name=?, attend=?, people=?, comment=?, days=?, updated=? WHERE id=?",
                      (name, attend, people, comment, json.dumps(days) if days is not None else None, now(), prev["id"]))
        else:
            c.execute("INSERT INTO rsvp(inv_id,guest_id,name,attend,people,comment,days,created,updated) VALUES(?,?,?,?,?,?,?,?,?)",
                      (r["id"], g["id"] if g else None, name, attend, people, comment, json.dumps(days) if days is not None else None, now(), now()))
    return JSONResponse({"ok": True})


async def pub_wish(req: Request):
    slug = req.path_params["slug"].lower()
    if not rate("wish:" + client_ip(req), 20, 3600):
        return err("Слишком много попыток. Попробуйте позже", 429)
    try:
        d = await body(req)
        idx = int(d.get("idx"))
    except Exception:
        return err("Неверный запрос")
    name = re.sub(r"\s+", " ", str(d.get("name", ""))).strip()[:80]
    if not name:
        return err("Укажите имя, чтобы молодожены знали, кто дарит")
    with db() as c:
        r = c.execute("SELECT * FROM invitations WHERE slug=? AND status='published'", (slug,)).fetchone()
        if not r:
            return err("Приглашение не найдено", 404)
        wl = json.loads(r["data"])["v"].get("wish", [])
        if not (0 <= idx < len(wl)):
            return err("Подарок не найден")
        try:
            c.execute("INSERT INTO wish VALUES(?,?,?,?)", (r["id"], idx, name, now()))
        except sqlite3.IntegrityError:
            return err("Этот подарок уже выбрал другой гость", 409)
    return JSONResponse({"ok": True})


# ---------------------------------------------------------------- cover pictures (link preview, stories)
OG_DIR = DATA / "og"
_RENDER_LOCK = threading.Lock()


def render_payload(d: dict, guest: str | None = None) -> dict:
    """What the cover needs: design, texts, greeting for one named guest. No guest list, photos only for covers."""
    v = {k: d["v"].get(k) for k in ("a", "b", "date", "time", "venue", "addr", "pal", "lat", "lng") if d["v"].get(k) is not None}
    v["photos"] = (d["v"].get("photos") or [])[:3]
    sel = [x for x in d["sel"] if x != "named"]
    if guest and "named" in d["sel"]:
        raw = [str(x) for x in d["v"].get("guests", [])]
        idx = next((i for i, x in enumerate(raw) if x.strip() == guest), -1)
        dear = d["v"].get("dear", [])
        v["guests"], v["dear"] = [guest], [dear[idx] if 0 <= idx < len(dear) else ""]
        sel.append("named")
    return {"tpl": d["tpl"], "sel": sel, "order": d["order"], "v": v}


RENDER_VER = "2"  # bump when the look of the pictures changes, so old cached pictures (and old preview links) are not reused


def render_hash(payload: dict, fmt: str) -> str:
    return hashlib.sha256((RENDER_VER + fmt + json.dumps(payload, sort_keys=True, ensure_ascii=False)).encode()).hexdigest()[:20]


_RENDER_JOBS: dict = {}


async def page_render(req: Request):
    """Cover for the screenshot; opened only by the server's own browser (one-time key)."""
    job = _RENDER_JOBS.pop(req.path_params["key"], None)
    h = guest_page()
    if not job or not h:
        return err("Не найдено", 404)
    data = json.dumps(job, ensure_ascii=False).replace("</", "<\\/")
    h = h.replace("<body", f"<script>window.RENDER={data};</script><body", 1)
    return HTMLResponse(h, headers={"Cache-Control": "no-store", "X-Robots-Tag": "noindex"})


def _render_chromium(payload: dict, fmt: str, out: Path) -> bool:
    try:
        from playwright.sync_api import sync_playwright
    except Exception:
        return False
    w, h = RENDER_SIZES[fmt]
    tmp = out.with_suffix(".tmp.png")
    try:
        with _RENDER_LOCK, sync_playwright() as pw:
            b = pw.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
            try:
                pg = b.new_page(viewport={"width": w, "height": h})
                key = secrets.token_urlsafe(16)
                _RENDER_JOBS[key] = {**payload, "fmt": fmt}
                pg.goto(f"http://127.0.0.1:{RENDER_PORT}/render/{key}", wait_until="load", timeout=15000)
                pg.wait_for_function("window.RENDER_DONE===true", timeout=10000)
                pg.screenshot(path=str(tmp), type="png")
            finally:
                b.close()
        tmp.replace(out)
        return True
    except Exception as e:
        print("render failed:", e, flush=True)
        return False


def _print_pdf(inv_id: str, session: str, who: int | None) -> bytes | None:
    """The printable version as a ready PDF (A5), made by the same Chromium from /print/<id> with the owner's session."""
    try:
        from playwright.sync_api import sync_playwright
    except Exception:
        return None
    try:
        with _RENDER_LOCK, sync_playwright() as pw:
            b = pw.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
            try:
                ctx = b.new_context(viewport={"width": 900, "height": 1200})
                ctx.add_cookies([{"name": "kn", "value": session, "url": f"http://127.0.0.1:{RENDER_PORT}"}])
                pg = ctx.new_page()
                pg.goto(f"http://127.0.0.1:{RENDER_PORT}/print/{inv_id}", wait_until="load", timeout=20000)
                pg.wait_for_selector(".sheet", timeout=15000)
                if who is not None and pg.query_selector("#who"):
                    pg.select_option("#who", str(who))
                pg.wait_for_function("document.fonts.status==='loaded' && [...document.images].every(i=>i.complete)", timeout=15000)
                pg.wait_for_timeout(300)
                return pg.pdf(prefer_css_page_size=True, print_background=True)
            finally:
                b.close()
    except Exception as e:
        print("print pdf failed:", e, flush=True)
        return None


def _print_zip(inv_id: str, session: str) -> bytes | None:
    """One PDF per named guest, all in a zip: easier than downloading 100 files one by one."""
    import zipfile
    try:
        from playwright.sync_api import sync_playwright
    except Exception:
        return None
    try:
        buf = io.BytesIO()
        with _RENDER_LOCK, sync_playwright() as pw, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            b = pw.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
            try:
                ctx = b.new_context(viewport={"width": 900, "height": 1200})
                ctx.add_cookies([{"name": "kn", "value": session, "url": f"http://127.0.0.1:{RENDER_PORT}"}])
                pg = ctx.new_page()
                pg.goto(f"http://127.0.0.1:{RENDER_PORT}/print/{inv_id}", wait_until="load", timeout=20000)
                pg.wait_for_selector(".sheet", timeout=15000)
                names = pg.eval_on_selector_all("#who option", "o=>o.map(x=>[x.value,x.textContent])") if pg.query_selector("#who") else []
                names = [(v, t) for v, t in names if v != "-1"] or [("-1", "Приглашение")]
                used = set()
                for k, (val, name) in enumerate(names, 1):
                    if val != "-1":
                        pg.select_option("#who", val)
                    pg.wait_for_function("document.fonts.status==='loaded' && [...document.images].every(i=>i.complete)", timeout=15000)
                    pg.wait_for_timeout(150)
                    fn = re.sub(r'[\\/:*?"<>|]+', " ", name).strip()[:60] or "Гость"
                    fn = f"{k:03d} {fn}.pdf"
                    if fn in used:
                        continue
                    used.add(fn)
                    z.writestr(fn, pg.pdf(prefer_css_page_size=True, print_background=True))
            finally:
                b.close()
        return buf.getvalue()
    except Exception as e:
        print("print zip failed:", e, flush=True)
        return None


@need_user
async def inv_print_zip(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        if "print" not in d["sel"]:
            return err("Печатная версия не подключена в этом приглашении", 403)
    data = await asyncio.to_thread(_print_zip, r["id"], req.cookies.get("kn", ""))
    if not data:
        return err("Не получилось собрать архив. Попробуйте еще раз через минуту", 503)
    fn = "knocky-" + (translit(inv_title(d)) or "priglasheniya") + ".zip"
    return Response(data, media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="{fn}"'})


@need_user
async def inv_print_pdf(req: Request):
    with db() as c:
        r = own_inv(c, req, req.path_params["id"])
        if not r:
            return err("Приглашение не найдено", 404)
        d = json.loads(r["data"])
        if "print" not in d["sel"]:
            return err("Печатная версия не подключена в этом приглашении", 403)
    try:
        who = int(req.query_params["i"]) if req.query_params.get("i", "") not in ("", "-1") else None
    except ValueError:
        who = None
    pdf = await asyncio.to_thread(_print_pdf, r["id"], req.cookies.get("kn", ""), who)
    if not pdf:
        return err("Не получилось собрать PDF. Нажмите «Распечатать» и выберите «Сохранить как PDF»", 503)
    fn = "knocky-" + (translit(inv_title(d)) or "priglashenie") + ".pdf"
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fn}"'})


def _render_card(payload: dict, fmt: str, out: Path) -> bool:
    """Fallback without a browser: a plain card in the design colours."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return False
    w, h = RENDER_SIZES[fmt]
    try:
        dz = json.loads((HERE / "static" / "assets" / "designs.json").read_text()).get(payload["tpl"], {})
    except Exception:
        dz = {}
    bg, ink, acc = dz.get("bg", "#F3EEEE"), dz.get("ink", "#241B20"), dz.get("acc", "#C62828")
    im = Image.new("RGB", (w, h), bg)
    dr = ImageDraw.Draw(im)
    def font(px, bold=False):
        for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
            try:
                return ImageFont.truetype(f, px)
            except Exception:
                pass
        return ImageFont.load_default()
    v = payload["v"]
    s = h / 630 if fmt == "og" else w / 900
    lines = []
    if v.get("guests"):
        dw = {"m": "Дорогой", "f": "Дорогая"}.get((v.get("dear") or [""])[0], "Дорогие")
        lines.append((f"{dw} {v['guests'][0]}!", font(int(30 * s)), acc))
    else:
        lines.append(("Приглашаем на свадьбу", font(int(30 * s)), acc))
    lines.append((f"{v.get('a') or ''} и {v.get('b') or ''}", font(int(78 * s), True), ink))
    p = (v.get("date") or "").split("-") if fmt != "og" else []
    months = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]
    if len(p) == 3 and p[1].isdigit() and 1 <= int(p[1]) <= 12:
        lines.append((f"{int(p[2])} {months[int(p[1]) - 1]} {p[0]}" + (f", {v['time']}" if v.get("time") else ""), font(int(34 * s)), ink))
    if v.get("venue") and fmt != "og":
        lines.append((v["venue"], font(int(28 * s)), ink))
    hs = [dr.textbbox((0, 0), t, font=f) for t, f, _ in lines]
    total = sum(b[3] - b[1] for b in hs) + int(26 * s) * (len(lines) - 1)
    y = (h - total) // 2
    for (t, f, c), b in zip(lines, hs):
        tw = b[2] - b[0]
        dr.text(((w - tw) // 2 - b[0], y - b[1]), t, font=f, fill=c)
        y += b[3] - b[1] + int(26 * s)
    tmp = out.with_suffix(".tmp.png")
    im.save(tmp, "PNG")
    tmp.replace(out)
    return True


def render_picture(payload: dict, fmt: str) -> Path | None:
    OG_DIR.mkdir(parents=True, exist_ok=True)
    out = OG_DIR / f"{fmt}-{render_hash(payload, fmt)}.png"
    if out.exists():
        return out
    if guest_page() and _render_chromium(payload, fmt, out):
        return out
    fb = OG_DIR / f"{fmt}-{render_hash(payload, fmt)}-card.png"
    if fb.exists() or _render_card(payload, fmt, fb):
        return fb
    return None


async def og_image(req: Request):
    """Link preview picture for knocky.ru/address (and for each named link: ?g=code)."""
    slug = req.path_params["slug"].lower()
    tok = req.query_params.get("g", "")[:20]
    with db() as c:
        r = c.execute("SELECT * FROM invitations WHERE slug=? AND status='published'", (slug,)).fetchone()
        if not r:
            return err("Не найдено", 404)
        guest = None
        if tok:
            g = c.execute("SELECT name FROM guests WHERE inv_id=? AND token=? AND active=1", (r["id"], tok)).fetchone()
            guest = g["name"] if g else None
    f = await asyncio.to_thread(render_picture, render_payload(json.loads(r["data"]), guest), "og")
    if not f:
        return err("Картинка недоступна", 503)
    return FileResponse(f, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})


# ---------------------------------------------------------------- pages
def guest_html() -> str | None:
    """The guest page contains the design code; it is stored encrypted and opened with the server key."""
    key = load_secrets().get("render_key")
    f = HERE / "sealed" / "guest.enc"
    if not key or not f.exists():
        return None
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    raw = f.read_bytes()
    try:
        return AESGCM(bytes.fromhex(key)).decrypt(raw[:12], raw[12:], None).decode()
    except Exception:
        return None


_GUEST = {"m": 0, "html": None}


def guest_page() -> str | None:
    f = HERE / "sealed" / "guest.enc"
    m = f.stat().st_mtime if f.exists() else 0
    if _GUEST["m"] != m or _GUEST["html"] is None:
        _GUEST["html"], _GUEST["m"] = guest_html(), m
    return _GUEST["html"]


SOON = """<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Knocky</title><style>body{margin:0;min-height:100vh;display:grid;place-items:center;background:#F3EEEE;color:#241B20;font:16px/1.5 system-ui,sans-serif;text-align:center;padding:24px}b{color:#C62828;font-size:22px}</style>
<div><b>knocky</b><p>Приглашение скоро будет доступно.</p></div></html>"""


def h_esc(s: str) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


async def page_guest(req: Request):
    h = guest_page()
    if not h:
        return HTMLResponse(SOON, headers={"Cache-Control": "no-cache", "X-Robots-Tag": "noindex"})
    title, desc, img = "Приглашение на свадьбу", "Приглашение сделано на Knocky", ""
    slug = req.path_params.get("slug")
    if slug:
        tok = req.query_params.get("g", "")[:20]
        with db() as c:
            r = c.execute("SELECT data, id FROM invitations WHERE slug=? AND status='published'", (slug.lower(),)).fetchone()
            g = c.execute("SELECT name FROM guests WHERE inv_id=? AND token=? AND active=1", (r["id"], tok)).fetchone() if r and tok else None
        if r:
            d = json.loads(r["data"])
            v = d["v"]
            ver = render_hash(render_payload(d, g["name"] if g else None), "og")[:10]
            img = f"{APP_URL}/og/{slug.lower()}.png?" + (f"g={tok}&" if g else "") + f"v={ver}"
            title = f"{inv_title({'v': v})} - приглашение на свадьбу"
            # no date or place in the preview text: messengers keep old previews, and the details may change
            desc = "Дата, место и программа - в приглашении"
    meta = (f'<title>{h_esc(title)}</title><meta property="og:title" content="{h_esc(title)}">'
            f'<meta property="og:description" content="{h_esc(desc)}"><meta name="description" content="{h_esc(desc)}">')
    if img:
        meta += (f'<meta property="og:image" content="{h_esc(img)}"><meta property="og:image:width" content="1200">'
                 f'<meta property="og:image:height" content="630"><meta name="twitter:card" content="summary_large_image">')
    h = h.replace("<title>Knocky</title>", meta, 1)
    return HTMLResponse(h, headers={"Cache-Control": "no-cache", "X-Robots-Tag": "noindex"})


async def page_cabinet(req: Request):
    return FileResponse(HERE / "static" / "index.html", headers={"Cache-Control": "no-cache"})


SETUP_FORM = """<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Knocky: настройка сервера</title><meta name="robots" content="noindex">
<style>body{margin:0;background:#F3EEEE;color:#241B20;font:16px/1.5 system-ui,-apple-system,sans-serif;padding:24px}
.c{max-width:460px;margin:6vh auto;background:#fff;border-radius:16px;padding:28px}h1{font-size:22px;margin:0 0 6px}h2{font-size:17px;margin:26px 0 0}
label{display:block;font-weight:600;margin:16px 0 6px}input{width:100%;box-sizing:border-box;font:inherit;padding:12px 14px;border:0;border-radius:12px;background:#F3EEEE}
button{margin-top:22px;width:100%;font:inherit;font-weight:700;padding:14px;border:0;border-radius:999px;background:#C62828;color:#fff;cursor:pointer}
.n{color:#6F6068;font-size:14px;margin:6px 0 0}.m{margin-top:16px;padding:12px 14px;border-radius:12px;background:#F3EEEE}</style>
<div class="c"><h1>Настройка сервера Knocky</h1><p class="n">Пароли сохранятся только на сервере. Их не видно ни в коде, ни в документах. Заполните одну часть или обе.</p>
__MSG__<form method="post"><input type="hidden" name="t" value="__T__">
<h2>Вход администратора без письма</h2>
<label for="ae">Почта администратора</label><input id="ae" name="admin_email" value="me@dudukchyan.ru" autocomplete="username">
<label for="ap">Придумайте пароль для входа</label><input id="ap" name="admin_pw" type="password" autocomplete="new-password" minlength="8" placeholder="не короче 8 символов">
<h2>Почта для писем с кодом</h2>
<label for="l">Почта, с которой уходят письма</label><input id="l" name="login" value="hello@knocky.ru">
<label for="p">Пароль приложения Яндекса для этой почты</label><input id="p" name="pw" type="password" autocomplete="off">
<label for="to">Куда отправить проверочное письмо</label><input id="to" name="to" value="me@dudukchyan.ru">
<h2>Оплата: Робокасса</h2>
<p class="n">Из личного кабинета Робокассы: Мои магазины - магазин Knocky - Технические настройки. Заполненные поля сохранятся, пустые останутся как были. __RK__</p>
<label for="rl">Идентификатор магазина</label><input id="rl" name="rk_login" autocomplete="off" data-lpignore="true" value="__RKL__">
<label for="r1">Пароль #1</label><input id="r1" name="rk_p1" type="password" autocomplete="new-password">
<label for="r2">Пароль #2</label><input id="r2" name="rk_p2" type="password" autocomplete="new-password">
<label for="t1">Тестовый пароль #1</label><input id="t1" name="rk_t1" type="password" autocomplete="new-password">
<label for="t2">Тестовый пароль #2</label><input id="t2" name="rk_t2" type="password" autocomplete="new-password">
<label for="sno">Система налогообложения ИП (для чеков)</label><select id="sno" name="rk_sno" style="width:100%;font:inherit;padding:12px 14px;border:0;border-radius:12px;background:#F3EEEE"><option value="usn_income"__S1__>УСН доходы</option><option value="usn_income_outcome"__S2__>УСН доходы минус расходы</option></select>
<label for="mode">Режим оплаты</label><select id="mode" name="rk_mode" style="width:100%;font:inherit;padding:12px 14px;border:0;border-radius:12px;background:#F3EEEE"><option value="test"__M1__>Тестовый: деньги не списываются</option><option value="live"__M2__>Боевой: настоящие оплаты</option></select>
<button>Сохранить</button></form></div></html>"""


def rk_conf_from(rk: dict) -> bool:
    test = rk.get("test", True)
    return bool(rk.get("login") and (rk.get("test_p1") and rk.get("test_p2") if test else rk.get("p1") and rk.get("p2")))


def setup_form(t: str, msg: str) -> str:
    rk = load_secrets().get("robokassa") or {}
    have = lambda k: "есть" if rk.get(k) else "нет"
    state = ("Сейчас: " + ("настроено, режим " + ("тестовый" if rk.get("test", True) else "боевой") if rk_conf_from(rk) else "не настроено")
             + f". Режим: {'тестовый' if rk.get('test', True) else 'боевой'}. Идентификатор: {h_esc(rk.get('login') or 'нет')}; пароль #1: {have('p1')}, #2: {have('p2')}; "
             f"тестовый #1: {have('test_p1')}, тестовый #2: {have('test_p2')}.")
    test = rk.get("test", True)
    sno2 = rk.get("sno") == "usn_income_outcome"
    return (SETUP_FORM.replace("__T__", t).replace("__MSG__", msg).replace("__RK__", state).replace("__RKL__", h_esc(rk.get("login", "")))
            .replace("__S1__", "" if sno2 else " selected").replace("__S2__", " selected" if sno2 else "")
            .replace("__M1__", " selected" if test else "").replace("__M2__", "" if test else " selected"))


def pw_hash(pw: str, salt: bytes | None = None) -> str:
    salt = salt or os.urandom(16)
    return salt.hex() + ":" + hashlib.pbkdf2_hmac("sha256", pw.encode(), salt, 300_000).hex()


def pw_ok(pw: str, stored: str) -> bool:
    try:
        salt, _ = stored.split(":")
        return hmac.compare_digest(pw_hash(pw, bytes.fromhex(salt)), stored)
    except Exception:
        return False


def setup_token_ok(t: str) -> bool:
    return bool(t) and SETUP_HASH != "SETUP_HASH_PLACEHOLDER" and hmac.compare_digest(sha(t), SETUP_HASH)


async def page_setup(req: Request):
    if req.method == "GET":
        t = req.query_params.get("t", "")
        if not setup_token_ok(t):
            return HTMLResponse("Ссылка недействительна", status_code=404)
        return HTMLResponse(setup_form(t, ""), headers={"Cache-Control": "no-store"})
    if not rate("setup:" + client_ip(req), 10, 3600):
        return HTMLResponse("Слишком много попыток", status_code=429)
    form = dict(await req.form())
    t = str(form.get("t", ""))
    if not setup_token_ok(t):
        return HTMLResponse("Ссылка недействительна", status_code=404)
    s = load_secrets()
    s["render_key"] = hashlib.sha256(("knocky-render:" + t).encode()).hexdigest()
    msgs = []
    ae, ap = str(form.get("admin_email") or "").strip().lower(), str(form.get("admin_pw") or "")
    if ap:
        if ae not in ADMINS:
            msgs.append(("bad", f"{h_esc(ae)} - не администратор. Администраторы: " + ", ".join(sorted(ADMINS))))
        elif len(ap) < 8:
            msgs.append(("bad", "Пароль администратора короче 8 символов, не сохранен"))
        else:
            s.setdefault("admin_pw", {})[ae] = pw_hash(ap)
            msgs.append(("ok", f"Пароль для входа {h_esc(ae)} сохранен. Входите на app.knocky.ru: «Вход для администратора по паролю»."))
    smtp = re.sub(r"\s+", "", str(form.get("pw") or ""))
    if smtp:
        trial = dict(s, smtp_login=str(form.get("login") or MAIL_FROM).strip(), smtp_password=smtp)
        to = str(form.get("to") or "me@dudukchyan.ru").strip()
        try:
            await asyncio.to_thread(send_mail, to, "Knocky: почта сервера работает",
                                    "Это проверочное письмо: сервер Knocky умеет отправлять письма.",
                                    mail_html("Почта работает", "Это проверочное письмо: сервер Knocky умеет отправлять письма. Можно входить в кабинет."), trial)
            s = trial
            msgs.append(("ok", f"Почта работает: проверочное письмо ушло на {h_esc(to)}."))
        except Exception as e:
            msgs.append(("bad", f"Письмо не ушло. Проверьте пароль приложения и что в настройках почты включен доступ для почтовых программ. ({type(e).__name__})"))
    rk = dict(s.get("robokassa") or {})
    changed = False
    for f, k in (("rk_login", "login"), ("rk_p1", "p1"), ("rk_p2", "p2"), ("rk_t1", "test_p1"), ("rk_t2", "test_p2")):
        val = str(form.get(f) or "").strip()
        if val and val != rk.get(k):
            rk[k], changed = val, True
    if form.get("rk_sno") in ("usn_income", "usn_income_outcome") and form.get("rk_sno") != rk.get("sno"):
        rk["sno"], changed = form["rk_sno"], True
    if form.get("rk_mode") in ("test", "live") and (form["rk_mode"] == "test") != rk.get("test", True):
        rk["test"], changed = form["rk_mode"] == "test", True
    if changed:
        s["robokassa"] = rk
        ready = rk_conf_from(rk)
        msgs.append(("ok" if ready else "bad", "Робокасса: настройки сохранены. " + ("Режим: " + ("тестовый." if rk.get("test", True) else "боевой, оплаты настоящие.") if ready else
                     "Не хватает: " + ", ".join(n for k, n in (("login", "идентификатора магазина"), ("test_p1" if rk.get("test", True) else "p1", "пароля #1" + (" (тестового)" if rk.get("test", True) else "")),
                     ("test_p2" if rk.get("test", True) else "p2", "пароля #2" + (" (тестового)" if rk.get("test", True) else ""))) if not rk.get(k)) + ". Оплата пока выключена.")))
    save_secrets(s)
    _GUEST["html"] = None
    if not msgs:
        msgs.append(("ok", "Сохранено."))
    msg = "".join(f'<div class="m"{" style=\"background:#FBE6E6\"" if k == "bad" else ""}>{m}</div>' for k, m in msgs)
    return HTMLResponse(setup_form(t, msg), headers={"Cache-Control": "no-store"})


async def health(req: Request):
    s = load_secrets()
    return JSONResponse({"ok": True, "mail": bool(s.get("smtp_password")), "guest_page": guest_page() is not None})


# ---------------------------------------------------------------- CORS / origin check
class Guard:
    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http":
            return await self.app(scope, receive, send)
        req = Request(scope)
        origin = req.headers.get("origin")
        path = scope["path"]
        if path.startswith("/api/"):
            if req.method == "OPTIONS":
                if origin in ORIGINS:
                    return await Response(status_code=204, headers=cors(origin) | {
                        "Access-Control-Allow-Methods": "GET,POST,PUT,DELETE", "Access-Control-Allow-Headers": "content-type",
                        "Access-Control-Max-Age": "600"})(scope, receive, send)
                return await Response(status_code=403)(scope, receive, send)
            if req.method not in ("GET", "HEAD") and origin and origin not in ORIGINS:
                return await err("Запрос с чужого сайта", 403)(scope, receive, send)

        async def send2(msg):
            if msg["type"] == "http.response.start":
                h = msg.setdefault("headers", [])
                if path.startswith("/api/") and origin in ORIGINS:
                    for k, v in cors(origin).items():
                        h.append((k.lower().encode(), v.encode()))
                h += [(b"x-content-type-options", b"nosniff"), (b"referrer-policy", b"strict-origin-when-cross-origin")]
                if not path.startswith("/i/"):
                    h.append((b"x-frame-options", b"SAMEORIGIN"))
            await send(msg)
        return await self.app(scope, receive, send2)


def cors(origin):
    return {"Access-Control-Allow-Origin": origin, "Access-Control-Allow-Credentials": "true", "Vary": "Origin"}


init_db()
routes = [
    Route("/api/health", health),
    Route("/api/auth/request", auth_request, methods=["POST"]),
    Route("/api/auth/verify", auth_verify, methods=["POST"]),
    Route("/api/auth/logout", auth_logout, methods=["POST"]),
    Route("/api/auth/password", auth_password, methods=["POST"]),
    Route("/api/me", me),
    Route("/api/drafts", draft_create, methods=["POST"]),
    Route("/api/drafts/claim", draft_claim, methods=["POST"]),
    Route("/api/invitations", inv_list),
    Route("/api/invitations", inv_create, methods=["POST"]),
    Route("/api/invitations/{id}", inv_get),
    Route("/api/invitations/{id}", inv_update, methods=["PUT"]),
    Route("/api/invitations/{id}", inv_delete, methods=["DELETE"]),
    Route("/api/invitations/{id}/publish", inv_publish, methods=["POST"]),
    Route("/api/invitations/{id}/unpublish", inv_unpublish, methods=["POST"]),
    Route("/api/invitations/{id}/answers", inv_answers),
    Route("/api/invitations/{id}/export.xlsx", inv_export),
    Route("/api/invitations/{id}/preview", preview_get),
    Route("/api/invitations/{id}/seating", inv_seating, methods=["PUT"]),
    Route("/api/invitations/{id}/price", inv_price),
    Route("/api/invitations/{id}/pay", inv_pay, methods=["POST"]),
    Route("/api/pay/result", pay_result, methods=["GET", "POST"]),
    Route("/pay/{kind:str}", pay_back),
    Route("/api/invitations/{id}/print.pdf", inv_print_pdf),
    Route("/api/invitations/{id}/print.zip", inv_print_zip),
    Route("/api/invitations/{id}/print", inv_print),
    Route("/api/public/{slug}", pub_get),
    Route("/api/public/{slug}/rsvp", pub_rsvp, methods=["POST"]),
    Route("/api/public/{slug}/wish", pub_wish, methods=["POST"]),
    Route("/setup", page_setup, methods=["GET", "POST"]),
    Route("/i/{slug}", page_guest),
    Route("/p/{id}", page_guest),
    Route("/print/{id}", page_guest),
    Route("/og/{slug}.png", og_image),
    Route("/render/{key}", page_render),
    Mount("/m", StaticFiles(directory=MEDIA, check_dir=False), name="media"),
    Mount("/assets", StaticFiles(directory=HERE / "static" / "assets", check_dir=False), name="assets"),
    Route("/", page_cabinet),
    Route("/inv/{id}", page_cabinet),
    Route("/admin", page_cabinet),
]
app = Guard(Starlette(routes=routes))
